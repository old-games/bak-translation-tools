"""Translation tables in the layout of the old-games.ru translation system.

Columns: A - original, B - translation (empty = keep the original), C - key
(where the text lives in the game: "DIAL_Z18:003422", "EXE:DS:0686"...), D - note
for translators. The format is chosen by the file extension:

  .xlsx - for the translation system: no header row (the system imports every
          row as a phrase); files exported from the system are read as is;
  .csv  - for tests and manual edits: UTF-8 with the header row HEADER.

Rows are matched by key, so their order and missing rows (phrases deleted in
the system) do not matter. The leading and trailing whitespace of a
translation is made the same as in the original (the paragraph tab of
dialogs...), whatever the translator typed at the edges.
"""

from __future__ import annotations

import csv
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADER = ("original", "translation", "key", "note")
FORMATS = ("xlsx", "csv")


@dataclass
class Row:
    key: str
    original: str
    translation: str = ""
    note: str = ""


def write_table(path: Path, rows: Iterable[Row]) -> int:
    """Write rows to .xlsx or .csv (by the extension); returns the row count."""
    rows = list(rows)
    cells = [(row.original, row.translation, row.key, row.note) for row in rows]
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with open(path, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(HEADER)
            writer.writerows(cells)
    elif suffix == ".xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = path.stem[:31]
        for values in cells:
            sheet.append(values)
        for line in sheet.iter_rows():
            for cell in line:
                cell.data_type = "s"  # never a formula, even if it starts with "="
        for column, width in zip("ABCD", (60, 60, 24, 40), strict=True):
            sheet.column_dimensions[column].width = width
        workbook.save(path)
    else:
        raise ValueError(f"{path}: unknown table format (expected .xlsx or .csv)")
    return len(rows)


def read_table(path: Path) -> list[Row]:
    """Rows of a .xlsx or .csv table; every row must have a unique key."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            lines = list(csv.reader(f))
        if lines and tuple(lines[0][:2]) == HEADER[:2]:
            lines = lines[1:]
    elif suffix == ".xlsx":
        # open the file ourselves: openpyxl checks the extension of a path
        with open(path, "rb") as f:
            sheet = load_workbook(f, read_only=True).active
            lines = [
                ["" if value is None else str(value) for value in line]
                for line in sheet.iter_rows(max_col=4, values_only=True)
            ]
    else:
        raise ValueError(f"{path}: unknown table format (expected .xlsx or .csv)")

    rows: list[Row] = []
    seen: set[str] = set()
    errors: list[str] = []
    for number, line in enumerate(lines, 1):
        original, translation, key, note = ([*line, "", "", "", ""])[:4]
        if not any((original, translation, key, note)):
            continue
        key = key.strip()
        if not key:
            errors.append(f"{path.name}, row {number}: no key (column C)")
        elif key in seen:
            errors.append(f"{path.name}, row {number}: duplicate key {key}")
        seen.add(key)
        rows.append(
            Row(key, _newlines(original), _newlines(translation), _newlines(note))
        )
    if errors:
        raise ValueError("\n".join(errors[:20]))
    return rows


def load_translations(path: Path) -> dict[str, tuple[str, str]]:
    """{key: (original, translation)}; the translation gets the edges of the
    original (empty translation = not translated)."""
    return {
        row.key: (row.original, fit_edges(row.original, row.translation))
        for row in read_table(path)
    }


def fit_edges(original: str, translation: str) -> str:
    """The translation with the leading/trailing whitespace of the original."""
    body = translation.strip()
    if not body:
        return ""
    core = original.strip()
    if not core:
        return translation
    start = original.index(core)
    return original[:start] + body + original[start + len(core) :]


def find_table(directory: Path, name: str) -> Path | None:
    """`name`.xlsx or `name`.csv in `directory` (None if neither; both = error)."""
    found = [p for p in (directory / f"{name}.{f}" for f in FORMATS) if p.exists()]
    if len(found) > 1:
        raise ValueError(f"both {found[0].name} and {found[1].name} in {directory}")
    return found[0] if found else None


def _newlines(text: str) -> str:
    return text.replace("\r\n", "\n")
