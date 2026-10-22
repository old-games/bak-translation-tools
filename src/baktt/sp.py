"""Exchange with the old-games.ru translation system ("Система Переводчик").

The system imports an XLS/XLSX sheet without a header: column A is the
original, column B the translation (empty = not translated); it trims both
cells and keeps no keys, only the row order. Its export has the same two
columns in the same order.

`sp export` turns the CSV files of `baktt export-all` into one XLSX per kind
of text plus a `.keys.csv` with the key and the full original of every row.
`sp import` reads an XLSX exported from the system, matches the rows by order
(checking the original) and writes the CSV back, restoring the leading and
trailing whitespace of the original (e.g. the paragraph tab of dialogs).
"""

from __future__ import annotations

import csv
from pathlib import Path

from cyclopts import App
from openpyxl import Workbook, load_workbook

from baktt.dialog import CSV_HEADER, load_translations

app = App(name="sp", help="Exchange with the old-games.ru translation system (XLSX)")

# CSV files of export-all that go to the system (DROPCAPS.csv is for artists)
TEXT_FILES = [
    "DDX.csv",
    "BOK.csv",
    "OBJINFO.csv",
    "KEYWORD.csv",
    "MNAMES.csv",
    "PARTY.csv",
    "UI.csv",
    "MISC.csv",
    "EXE.csv",
]
KEYS_HEADER = ("row", "key", "english")


def _read_rows(csv_path: Path) -> list[tuple[str, str, str]]:
    """(key, english, russian) of a CSV in the common format or of BOK.csv
    (sections "===NAME.BOK" followed by rows "english,russian")."""
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    if csv_path.name.upper() == "BOK.CSV":
        result = []
        section = ""
        index = 0
        for row in rows:
            if not row:
                continue
            if row[0].startswith("==="):
                section, index = row[0][3:], 0
                continue
            result.append(
                (f"{section}:{index:03d}", row[0], row[1] if len(row) > 1 else "")
            )
            index += 1
        return result
    translations = load_translations(csv_path)
    return [(key, english, russian) for key, (english, russian) in translations.items()]


def _normalize(text: str) -> str:
    return text.replace("\r\n", "\n").strip()


@app.command(name="export")
def export_command(src_dir: Path, out_dir: Path) -> None:
    """Make XLSX files for the translation system from the export-all CSVs."""
    sp_export(src_dir, out_dir)


def sp_export(src_dir: Path, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    for name in TEXT_FILES:
        csv_path = src_dir / name
        if not csv_path.exists():
            continue
        rows = _read_rows(csv_path)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = Path(name).stem
        for i, (_key, english, russian) in enumerate(rows, 1):
            for column, value in (("A", english.strip()), ("B", russian.strip())):
                cell = sheet[f"{column}{i}"]
                cell.value = value
                cell.data_type = "s"  # never a formula, even if it starts with "="
        sheet.column_dimensions["A"].width = 60
        sheet.column_dimensions["B"].width = 60
        stem = Path(name).stem
        workbook.save(out_dir / f"{stem}.xlsx")
        with open(out_dir / f"{stem}.keys.csv", "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(KEYS_HEADER)
            for i, (key, english, _russian) in enumerate(rows, 1):
                writer.writerow((i, key, english))
        print(f"{stem}.xlsx: {len(rows)} rows")


@app.command(name="import")
def import_command(xlsx_path: Path, keys_path: Path, csv_path: Path) -> None:
    """Convert an XLSX exported from the translation system back to CSV.

    `keys_path` is the .keys.csv made by `sp export` for this file; rows are
    matched by order and the original text is checked.
    """
    sp_import(xlsx_path, keys_path, csv_path)


def sp_import(xlsx_path: Path, keys_path: Path, csv_path: Path) -> None:
    with open(keys_path, encoding="utf-8-sig", newline="") as f:
        keys = list(csv.reader(f))[1:]
    sheet = load_workbook(xlsx_path, read_only=True).active
    cells = [
        ("" if a is None else str(a), "" if b is None else str(b))
        for a, b, *_ in sheet.iter_rows(min_col=1, max_col=2, values_only=True)
    ]
    while cells and cells[-1] == ("", ""):
        cells.pop()
    errors = []
    if len(cells) != len(keys):
        errors.append(f"{xlsx_path}: {len(cells)} rows, {keys_path} has {len(keys)}")
    rows = []
    for (row, key, english), (original, russian) in zip(keys, cells, strict=False):
        if _normalize(original) != _normalize(english):
            errors.append(f"row {row} ({key}): the original differs from the game text")
            continue
        russian = russian.replace("\r\n", "\n").strip()
        if russian:
            # restore the whitespace the system trims (paragraph tab, newlines)
            body = english.strip()
            start = english.index(body[:1]) if body else 0
            russian = english[:start] + russian + english[start + len(body) :]
        rows.append((key, english, russian))
    for error in errors[:20]:
        print("ERROR", error)
    if errors:
        raise SystemExit(f"{len(errors)} errors")

    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with open(csv_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        if keys_path.name.upper().startswith("BOK."):
            section = None
            for key, english, russian in rows:
                name = key.rsplit(":", 1)[0]
                if name != section:
                    writer.writerow((f"==={name}",))
                    section = name
                writer.writerow((english, russian))
        else:
            writer.writerow(CSV_HEADER)
            writer.writerows(rows)
    translated = sum(1 for row in rows if row[2])
    print(f"{csv_path}: {len(rows)} rows, {translated} translated")
